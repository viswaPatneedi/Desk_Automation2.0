#!/usr/bin/env python3
"""
Validate Results Method Implementation
Execute SSH commands and validate output against user-provided expected results
"""

import sys
import time
import os
import re
import paramiko
from datetime import datetime, timezone

# Import shared utilities
from methods.method_utils import (
    log_message,
    fetch_build_details,
    create_screenshot_folder,
    create_execution_log_path,
    get_folder_method_name,
    get_lexar_base_path
)

# Import screenshot utilities
from utils.screenshot_utils import take_and_analyze_screenshot
from tools.screen.screenshot_utils_vnc import take_vnc_screenshot_with_fallback

def _parse_top_number(value):
    try:
        v = value.strip().lower()
        if v.endswith('g'):
            return float(v[:-1]) * 1024 * 1024
        if v.endswith('m'):
            return float(v[:-1]) * 1024
        if v.endswith('k'):
            return float(v[:-1])
        return float(v)
    except Exception:
        return None

def _parse_top_output(output_text):
    time_value = None
    mem_total = mem_free = mem_used = buff_cache = None
    swap_total = swap_free = swap_used = avail_mem = None
    pid = virt = res = shr = cpu_pct = mem_pct = None

    lines = output_text.splitlines()
    header_index = {}

    for line in lines:
        if line.startswith('top -'):
            # Example: top - 04:29:03 up ...
            parts = line.split()
            if len(parts) >= 3:
                time_value = parts[2]
        if 'Mem' in line and 'total' in line and 'free' in line:
            # Example: MiB Mem : 3604.6 total, 2236.3 free, 895.8 used, 472.5 buff/cache
            try:
                numbers = [p.strip().replace(',', '') for p in line.replace(':', ' ').split()]
                if 'total' in numbers:
                    idx = numbers.index('total')
                    mem_total = _parse_top_number(numbers[idx - 1])
                if 'free' in numbers:
                    idx = numbers.index('free')
                    mem_free = _parse_top_number(numbers[idx - 1])
                if 'used' in numbers:
                    idx = numbers.index('used')
                    mem_used = _parse_top_number(numbers[idx - 1])
                if 'buff/cache' in numbers:
                    idx = numbers.index('buff/cache')
                    buff_cache = _parse_top_number(numbers[idx - 1])
            except Exception:
                pass
        if 'Swap' in line and 'total' in line and 'free' in line:
            # Example: MiB Swap: 1802.3 total, 1802.3 free, 0.0 used. 2578.1 avail Mem
            try:
                tokens = [p.strip().replace(',', '') for p in line.replace(':', ' ').replace('.', ' ').split()]
                if 'total' in tokens:
                    idx = tokens.index('total')
                    swap_total = _parse_top_number(tokens[idx - 1])
                if 'free' in tokens:
                    idx = tokens.index('free')
                    swap_free = _parse_top_number(tokens[idx - 1])
                if 'used' in tokens:
                    idx = tokens.index('used')
                    swap_used = _parse_top_number(tokens[idx - 1])
                if 'avail' in tokens and 'Mem' in tokens:
                    idx = tokens.index('avail')
                    avail_mem = _parse_top_number(tokens[idx - 1])
            except Exception:
                pass
        if line.strip().startswith('PID '):
            header_cols = line.split()
            header_index = {name: idx for idx, name in enumerate(header_cols)}

    for line in lines:
        if 'WPEFramework' in line:
            parts = line.split()
            if header_index:
                def get_col(name):
                    idx = header_index.get(name)
                    if idx is not None and idx < len(parts):
                        return parts[idx]
                    return None
                pid = get_col('PID')
                virt = _parse_top_number(get_col('VIRT') or '')
                res = _parse_top_number(get_col('RES') or '')
                shr = _parse_top_number(get_col('SHR') or '')
                cpu_pct = _parse_top_number(get_col('%CPU') or '')
                mem_pct = _parse_top_number(get_col('%MEM') or '')
            break

    return {
        "time": time_value,
        "mem_total": mem_total,
        "mem_free": mem_free,
        "mem_used": mem_used,
        "buff_cache": buff_cache,
        "swap_total": swap_total,
        "swap_free": swap_free,
        "swap_used": swap_used,
        "avail_mem": avail_mem,
        "pid": pid,
        "virt": virt,
        "res": res,
        "shr": shr,
        "cpu_pct": cpu_pct,
        "mem_pct": mem_pct
    }

def _write_top_output_to_excel(device_ip, device_name, iteration, command, samples,
                               expected_output, validation_type, validation_passed,
                               section_label=None, combined_method_name=None):
    try:
        from openpyxl import Workbook, load_workbook
    except Exception as exc:
        log_message(f"⚠ Could not load Excel library (openpyxl): {exc}")
        return

    try:
        # Build USB output path
        current_date = datetime.now(timezone.utc).strftime('%Y-%m-%d')
        safe_device_name = device_name.replace(' ', '-').replace('/', '_').replace('\\', '_').replace('__', '_').replace('--', '-').upper()
        safe_ip = device_ip.replace('.', '-')
        device_folder = f"{safe_ip}_{safe_device_name}"
        output_dir = os.path.join(get_lexar_base_path(), "EXECUTION_LOGS", current_date, device_folder)
        os.makedirs(output_dir, exist_ok=True)

        method_tag = (get_folder_method_name() or combined_method_name or "validate_results").replace(',', '_')
        file_name = f"{safe_ip}_{safe_device_name}_VALIDATE_RESULTS_{method_tag}_TOP_OUTPUT.xlsx"
        file_path = os.path.join(output_dir, file_name)

        if os.path.exists(file_path):
            workbook = load_workbook(file_path)
        else:
            workbook = Workbook()
            # Remove default sheet
            if workbook.sheetnames:
                default_sheet = workbook[workbook.sheetnames[0]]
                workbook.remove(default_sheet)

        sheet_name = f"ITR-{iteration}"
        if sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
        else:
            ws = workbook.create_sheet(title=sheet_name)
            ws.append(["Timestamp UTC", datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')])
            ws.append(["Device IP", device_ip])
            ws.append(["Device Name", device_name])
            ws.append(["Iteration", iteration])
            ws.append([])

        label = section_label or "TOP Capture"
        ws.append([label])
        ws.append([
            "Time",
            "MemTotal",
            "MemFree",
            "MemUsed",
            "BuffCache",
            "SwapTotal",
            "SwapFree",
            "SwapUsed",
            "AvailMem",
            "RSS_KB",
            "PSS_KB",
            "PID",
            "VIRT_KiB",
            "RES_KiB",
            "SHR_KiB",
            "CPU_pct",
            "MEM_pct"
        ])

        totals = {
            "mem_total": 0.0, "mem_free": 0.0, "mem_used": 0.0, "buff_cache": 0.0,
            "swap_total": 0.0, "swap_free": 0.0, "swap_used": 0.0, "avail_mem": 0.0,
            "virt": 0.0, "res": 0.0, "shr": 0.0, "cpu_pct": 0.0, "mem_pct": 0.0
        }
        count = 0

        for sample in samples:
            if not sample:
                continue
            count += 1
            ws.append([
                sample.get("time"),
                sample.get("mem_total"),
                sample.get("mem_free"),
                sample.get("mem_used"),
                sample.get("buff_cache"),
                sample.get("swap_total"),
                sample.get("swap_free"),
                sample.get("swap_used"),
                sample.get("avail_mem"),
                sample.get("res"),
                None,
                sample.get("pid"),
                sample.get("virt"),
                sample.get("res"),
                sample.get("shr"),
                sample.get("cpu_pct"),
                sample.get("mem_pct")
            ])

            for key in totals:
                val = sample.get(key)
                if val is not None:
                    totals[key] += val

        if count > 0:
            ws.append([
                "AVERAGE",
                totals["mem_total"] / count,
                totals["mem_free"] / count,
                totals["mem_used"] / count,
                totals["buff_cache"] / count,
                totals["swap_total"] / count,
                totals["swap_free"] / count,
                totals["swap_used"] / count,
                totals["avail_mem"] / count,
                totals["res"] / count,
                None,
                None,
                totals["virt"] / count,
                totals["res"] / count,
                totals["shr"] / count,
                totals["cpu_pct"] / count,
                totals["mem_pct"] / count
            ])
        else:
            ws.append(["AVERAGE", "(no samples)"])

        ws.append([])

        workbook.save(file_path)
        log_message(f"💾 TOP output saved to Excel: {file_path}")
    except Exception as exc:
        log_message(f"⚠ Could not write TOP output to Excel: {exc}")

def validate_command_output(device_ip, port, username, password, iteration=1, device_name="Device", 
                            command=None, expected_output=None, validation_type="contains", 
                            combined_method_name=None):
    """
    Execute SSH command and validate output against expected results
    
    Args:
        device_ip: Device IP address
        port: SSH port
        username: SSH username
        password: SSH password
        iteration: Current iteration number
        device_name: Display name of device
        command: Shell command to execute on device
        expected_output: Expected output text or regex pattern (for validation)
        validation_type: Type of validation:
            - "contains" (default): Check if expected_output is a substring of command output
            - "exact": Check for exact match with command output
            - "not_contains": Check that expected_output is NOT in command output
            - "regex": Treat expected_output as a regex pattern and validation passes if pattern matches
            - "not_regex": Treat expected_output as a regex pattern and validation passes if pattern does NOT match (useful for multiple patterns with alternation)
        combined_method_name: Name for combined method execution (optional)
    
    Returns:
        Dict with iteration, screenshots, logs, success status, and validation details
    """
    timestamp = datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')
    screenshots_list = []
    logs_list = []
    
    # Clean device name for filename
    safe_device_name = device_name.replace(' ', '_').replace('/', '_').replace('\\', '_')
    
    # Set USB log file path for this execution
    try:
        from services.log_service import LogService
        log_service = LogService()
        usb_log_path = create_execution_log_path(device_ip, device_name, iteration, "VALIDATE_RESULTS")
        log_service.current_usb_log_file = usb_log_path
        log_message(f"💾 USB execution log: {usb_log_path}")
    except Exception as e:
        log_message(f"⚠ Could not set USB log path: {e}")
    
    log_message("="*80)
    log_message("VALIDATE RESULTS - START")
    log_message("="*80)
    
    # Validate required parameters
    if not command:
        log_message("❌ ERROR: No command provided for execution")
        return {"iteration": iteration, "screenshots": screenshots_list, "logs": logs_list, "success": False, "details": "No command provided"}
    
    if not expected_output:
        log_message("❌ ERROR: No expected output provided for validation")
        return {"iteration": iteration, "screenshots": screenshots_list, "logs": logs_list, "success": False, "details": "No expected output provided"}
    
    log_message(f"📋 Command: {command}")
    log_message(f"🎯 Expected Output: {expected_output}")
    log_message(f"🔍 Validation Type: {validation_type.upper()}")
    
    try:
        # STEP 1: CONNECT TO DEVICE
        log_message("\n[STEP 1] Connecting to device...")
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        ssh.connect(device_ip, port=port, username=username, password=password, timeout=15)
        log_message("✓ Connected to device successfully")
        
        # Fetch build details from device
        fetch_build_details(ssh, log_message)
        
        # STEP 2: EXECUTE COMMAND
        command_to_run = command
        section_label = None
        duration_seconds = 300
        interval_seconds = 10
        if '#' in command:
            command_to_run, label = command.split('#', 1)
            command_to_run = command_to_run.strip()
            section_label = label.strip() or None
            # Optional overrides in label: duration=300 interval=10
            if section_label:
                lower_label = section_label.lower()
                try:
                    if 'duration=' in lower_label:
                        duration_seconds = int(lower_label.split('duration=')[1].split()[0])
                    if 'interval=' in lower_label:
                        interval_seconds = int(lower_label.split('interval=')[1].split()[0])
                except Exception:
                    pass

        is_top_command = command_to_run.strip().lower().startswith("top")
        log_message(f"\n[STEP 2] Executing command: {command_to_run}")

        command_output = ""
        command_error = ""
        exit_status = 0
        top_samples = []

        if is_top_command:
            log_message(f"⏱ Collecting TOP samples for {duration_seconds}s every {interval_seconds}s...")
            end_time = time.time() + max(duration_seconds, 1)
            while time.time() < end_time:
                stdin, stdout, stderr = ssh.exec_command(command_to_run, timeout=30)
                sample_output = stdout.read().decode('utf-8', errors='ignore').strip()
                sample_error = stderr.read().decode('utf-8', errors='ignore').strip()
                exit_status = stdout.channel.recv_exit_status()
                if sample_error:
                    command_error = sample_error
                command_output = sample_output
                parsed = _parse_top_output(sample_output)
                if parsed:
                    top_samples.append(parsed)
                time.sleep(max(interval_seconds, 1))

            log_message(f"✓ TOP collection complete ({len(top_samples)} samples)")
        else:
            stdin, stdout, stderr = ssh.exec_command(command_to_run, timeout=30)
            command_output = stdout.read().decode('utf-8', errors='ignore').strip()
            command_error = stderr.read().decode('utf-8', errors='ignore').strip()
            exit_status = stdout.channel.recv_exit_status()
            log_message(f"✓ Command executed (Exit Status: {exit_status})")
        
        # Display output
        if command_output:
            log_message(f"\n📤 Command Output ({len(command_output)} characters):")
            log_message("-" * 80)
            # Limit output display to first 1000 characters for log readability
            display_output = command_output[:1000]
            if len(command_output) > 1000:
                display_output += f"\n... (truncated, showing first 1000 of {len(command_output)} characters)"
            log_message(display_output)
            log_message("-" * 80)
        else:
            log_message("📤 Command Output: (empty)")
        
        if command_error:
            log_message(f"\n⚠ Command Error Output:")
            log_message(command_error[:500])
        
        # STEP 3: VALIDATE OUTPUT
        log_message(f"\n[STEP 3] Validating output...")
        validation_passed = False
        validation_details = ""
        
        if validation_type == "contains":
            # Check if expected output is contained in command output
            if expected_output in command_output:
                validation_passed = True
                validation_details = f"Output CONTAINS expected text: '{expected_output}'"
                log_message(f"✓ {validation_details}")
            else:
                validation_passed = False
                validation_details = f"Output does NOT contain expected text: '{expected_output}'"
                log_message(f"❌ {validation_details}")
        
        elif validation_type == "exact":
            # Check for exact match
            if command_output == expected_output:
                validation_passed = True
                validation_details = f"Output EXACTLY matches expected text"
                log_message(f"✓ {validation_details}")
            else:
                validation_passed = False
                validation_details = f"Output does NOT match expected text exactly"
                log_message(f"❌ {validation_details}")
                log_message(f"   Expected: '{expected_output}'")
                log_message(f"   Got:      '{command_output[:200]}{'...' if len(command_output) > 200 else ''}'")
        
        elif validation_type == "not_contains":
            # Check that output does NOT contain the text
            if expected_output not in command_output:
                validation_passed = True
                validation_details = f"Output correctly does NOT contain: '{expected_output}'"
                log_message(f"✓ {validation_details}")
            else:
                validation_passed = False
                validation_details = f"Output incorrectly CONTAINS: '{expected_output}'"
                log_message(f"❌ {validation_details}")
        
        elif validation_type == "regex":
            # Check if regex pattern matches in command output
            try:
                if re.search(expected_output, command_output):
                    validation_passed = True
                    validation_details = f"Output MATCHES regex pattern: '{expected_output}'"
                    log_message(f"✓ {validation_details}")
                else:
                    validation_passed = False
                    validation_details = f"Output does NOT match regex pattern: '{expected_output}'"
                    log_message(f"❌ {validation_details}")
            except re.error as regex_err:
                validation_passed = False
                validation_details = f"Invalid regex pattern: {regex_err}"
                log_message(f"❌ {validation_details}")
        
        elif validation_type == "not_regex":
            # Check that regex pattern does NOT match in command output (inverse of regex)
            # Useful for validating that unwanted patterns are absent
            try:
                if re.search(expected_output, command_output):
                    validation_passed = False
                    validation_details = f"Validation FAILED: Output MATCHES unwanted regex pattern: '{expected_output}'"
                    log_message(f"❌ {validation_details}")
                else:
                    validation_passed = True
                    validation_details = f"Validation PASSED: Output does NOT contain unwanted pattern: '{expected_output}'"
                    log_message(f"✓ {validation_details}")
            except re.error as regex_err:
                validation_passed = False
                validation_details = f"Invalid regex pattern: {regex_err}"
                log_message(f"❌ {validation_details}")
        
        else:
            validation_passed = False
            validation_details = f"Unknown validation type: '{validation_type}'"
            log_message(f"❌ {validation_details}")

        # Save TOP output to Excel (USB) if command is TOP
        if is_top_command:
            _write_top_output_to_excel(
                device_ip=device_ip,
                device_name=device_name,
                iteration=iteration,
                command=command_to_run,
                samples=top_samples,
                expected_output=expected_output,
                validation_type=validation_type,
                validation_passed=validation_passed,
                section_label=section_label,
                combined_method_name=combined_method_name
            )
        
        # STEP 4: CAPTURE SCREENSHOT (OPTIONAL)
        log_message(f"\n[STEP 4] Capturing screenshot...")
        method_for_folder = get_folder_method_name() or combined_method_name or "validate_results"
        screenshot_folder = create_screenshot_folder(device_ip, safe_device_name, iteration, "After", method_for_folder)
        
        screenshot_status = "SUCCESS" if validation_passed else "FAILED"
        screenshot_name = f"{device_ip}_{safe_device_name}_Iteration-{iteration}_Validate-Results-{screenshot_status}_{timestamp}"
        
        try:
            # Activate ScreenCapture service
            from methods.method_utils import activate_screencapture_service
            activate_screencapture_service(ssh, log_message)
            
            screenshot_result = take_vnc_screenshot_with_fallback(
                    ssh=ssh,
                    device_ip=device_ip,
                    device_name=safe_device_name,
                    iteration=iteration,
                    screenshot_folder=screenshot_folder,
                    log_callback=log_message,
                    fallback_to_plugin=True,
                    context="Default"
                )
            if screenshot_result and screenshot_result.get('success'):
                log_message(f"✓ Screenshot saved: {screenshot_result.get('local_path')}")
                screenshots_list.append(screenshot_result.get('local_path', ''))
        except Exception as screenshot_error:
            log_message(f"⚠ Could not capture screenshot: {screenshot_error}")
        
        # Close SSH connection
        ssh.close()
        
        # FINAL RESULT
        log_message("\n" + "="*80)
        if validation_passed:
            log_message("✓ VALIDATE RESULTS - PASSED")
            log_message(f"✓ {validation_details}")
        else:
            log_message("❌ VALIDATE RESULTS - FAILED")
            log_message(f"❌ {validation_details}")
        log_message("="*80)
        
        return {
            "iteration": iteration,
            "screenshots": screenshots_list,
            "logs": logs_list,
            "success": validation_passed,
            "details": validation_details,
            "command": command,
            "output": command_output[:500],  # Store first 500 chars of output
            "expected": expected_output,
            "validation_type": validation_type
        }
    
    except paramiko.SSHException as ssh_error:
        log_message(f"❌ SSH Error: {ssh_error}")
        return {
            "iteration": iteration,
            "screenshots": screenshots_list,
            "logs": logs_list,
            "success": False,
            "details": f"SSH connection failed: {str(ssh_error)}"
        }
    
    except Exception as e:
        import traceback
        log_message(f"❌ Error during Validate Results execution: {e}")
        log_message(f"Full traceback:\n{traceback.format_exc()}")
        return {
            "iteration": iteration,
            "screenshots": screenshots_list,
            "logs": logs_list,
            "success": False,
            "details": f"Execution error: {str(e)}"
        }
