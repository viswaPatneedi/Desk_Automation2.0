"""
System Command Method
Execute approved/system commands on device and optionally capture TOP metrics to Excel.
"""

import time
import os
import paramiko
from datetime import datetime, timezone

from method_utils import (
    log_message,
    get_folder_method_name,
    get_lexar_base_path
)


def _get_system_command_log_path(device_ip, device_name, iteration, method_folder=None):
    current_date = datetime.now(timezone.utc).strftime('%Y-%m-%d')
    safe_device_name = device_name.replace(' ', '-').replace('/', '_').replace('\\', '_').replace('__', '_').replace('--', '-').upper()
    safe_ip = device_ip.replace('.', '-')
    device_folder = f"{safe_ip}_{safe_device_name}"
    folder_name = (method_folder or "SYSTEM_COMMAND").replace(' ', '_')
    itr_dir = os.path.join(get_lexar_base_path(), "EXECUTION_LOGS", current_date, device_folder, folder_name, f"ITR-{iteration}")
    os.makedirs(itr_dir, exist_ok=True)
    return os.path.join(itr_dir, f"SYSTEM_COMMAND_ITR_{iteration}.log")


def _parse_top_number(value):
    try:
        v = value.strip().lower()
        if v.endswith('g'):
            return float(v[:-1]) * 1024 * 1024
        if v.endswith('m'):
            return float(v[:-1]) * 1024
        if v.endswith('k'):
            return float(v[:-1])
        return float(v)
    except Exception:
        return None


def _parse_top_output(output_text):
    time_value = None
    mem_total = mem_free = mem_used = buff_cache = None
    swap_total = swap_free = swap_used = avail_mem = None
    pid = virt = res = shr = cpu_pct = mem_pct = None

    lines = output_text.splitlines()
    header_index = {}

    for line in lines:
        if line.startswith('top -'):
            parts = line.split()
            if len(parts) >= 3:
                time_value = parts[2]
        if 'Mem' in line and 'total' in line and 'free' in line:
            try:
                numbers = [p.strip().replace(',', '') for p in line.replace(':', ' ').split()]
                if 'total' in numbers:
                    idx = numbers.index('total')
                    mem_total = _parse_top_number(numbers[idx - 1])
                if 'free' in numbers:
                    idx = numbers.index('free')
                    mem_free = _parse_top_number(numbers[idx - 1])
                if 'used' in numbers:
                    idx = numbers.index('used')
                    mem_used = _parse_top_number(numbers[idx - 1])
                if 'buff/cache' in numbers:
                    idx = numbers.index('buff/cache')
                    buff_cache = _parse_top_number(numbers[idx - 1])
            except Exception:
                pass
        if 'Swap' in line and 'total' in line and 'free' in line:
            try:
                tokens = [p.strip().replace(',', '') for p in line.replace(':', ' ').replace('.', ' ').split()]
                if 'total' in tokens:
                    idx = tokens.index('total')
                    swap_total = _parse_top_number(tokens[idx - 1])
                if 'free' in tokens:
                    idx = tokens.index('free')
                    swap_free = _parse_top_number(tokens[idx - 1])
                if 'used' in tokens:
                    idx = tokens.index('used')
                    swap_used = _parse_top_number(tokens[idx - 1])
                if 'avail' in tokens and 'Mem' in tokens:
                    idx = tokens.index('avail')
                    avail_mem = _parse_top_number(tokens[idx - 1])
            except Exception:
                pass
        if line.strip().startswith('PID '):
            header_cols = line.split()
            header_index = {name: idx for idx, name in enumerate(header_cols)}

    for line in lines:
        if 'WPEFramework' in line:
            parts = line.split()
            if header_index:
                def get_col(name):
                    idx = header_index.get(name)
                    if idx is not None and idx < len(parts):
                        return parts[idx]
                    return None
                pid = get_col('PID')
                virt = _parse_top_number(get_col('VIRT') or '')
                res = _parse_top_number(get_col('RES') or '')
                shr = _parse_top_number(get_col('SHR') or '')
                cpu_pct = _parse_top_number(get_col('%CPU') or '')
                mem_pct = _parse_top_number(get_col('%MEM') or '')
            break

    return {
        "time": time_value,
        "mem_total": mem_total,
        "mem_free": mem_free,
        "mem_used": mem_used,
        "buff_cache": buff_cache,
        "swap_total": swap_total,
        "swap_free": swap_free,
        "swap_used": swap_used,
        "avail_mem": avail_mem,
        "pid": pid,
        "virt": virt,
        "res": res,
        "shr": shr,
        "cpu_pct": cpu_pct,
        "mem_pct": mem_pct
    }


def _write_top_output_to_excel(device_ip, device_name, iteration, command, samples,
                               section_label=None, combined_method_name=None):
    try:
        from openpyxl import Workbook, load_workbook
    except Exception as exc:
        log_message(f"⚠ Could not load Excel library (openpyxl): {exc}")
        return

    try:
        current_date = datetime.now(timezone.utc).strftime('%Y-%m-%d')
        safe_device_name = device_name.replace(' ', '-').replace('/', '_').replace('\\', '_').replace('__', '_').replace('--', '-').upper()
        safe_ip = device_ip.replace('.', '-')
        device_folder = f"{safe_ip}_{safe_device_name}"
        output_dir = os.path.join(get_lexar_base_path(), "EXECUTION_LOGS", current_date, device_folder)
        os.makedirs(output_dir, exist_ok=True)

        method_tag = (get_folder_method_name() or combined_method_name or "system_command").replace(',', '_')
        file_name = f"{safe_ip}_{safe_device_name}_SYSTEM_COMMAND_{method_tag}_TOP_OUTPUT.xlsx"
        file_path = os.path.join(output_dir, file_name)

        if os.path.exists(file_path):
            workbook = load_workbook(file_path)
        else:
            workbook = Workbook()
            if workbook.sheetnames:
                default_sheet = workbook[workbook.sheetnames[0]]
                workbook.remove(default_sheet)

        sheet_name = f"ITR-{iteration}"
        if sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
        else:
            ws = workbook.create_sheet(title=sheet_name)
            ws.append(["Timestamp UTC", datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')])
            ws.append(["Device IP", device_ip])
            ws.append(["Device Name", device_name])
            ws.append(["Iteration", iteration])
            ws.append([])

        label = section_label or "TOP Capture"
        ws.append([label])
        ws.append([
            "Time",
            "MemTotal",
            "MemFree",
            "MemUsed",
            "BuffCache",
            "SwapTotal",
            "SwapFree",
            "SwapUsed",
            "AvailMem",
            "RSS_KB",
            "PSS_KB",
            "PID",
            "VIRT_KiB",
            "RES_KiB",
            "SHR_KiB",
            "CPU_pct",
            "MEM_pct"
        ])

        totals = {
            "mem_total": 0.0, "mem_free": 0.0, "mem_used": 0.0, "buff_cache": 0.0,
            "swap_total": 0.0, "swap_free": 0.0, "swap_used": 0.0, "avail_mem": 0.0,
            "virt": 0.0, "res": 0.0, "shr": 0.0, "cpu_pct": 0.0, "mem_pct": 0.0
        }
        count = 0

        for sample in samples:
            if not sample:
                continue
            count += 1
            ws.append([
                sample.get("time"),
                sample.get("mem_total"),
                sample.get("mem_free"),
                sample.get("mem_used"),
                sample.get("buff_cache"),
                sample.get("swap_total"),
                sample.get("swap_free"),
                sample.get("swap_used"),
                sample.get("avail_mem"),
                sample.get("res"),
                None,
                sample.get("pid"),
                sample.get("virt"),
                sample.get("res"),
                sample.get("shr"),
                sample.get("cpu_pct"),
                sample.get("mem_pct")
            ])

            for key in totals:
                val = sample.get(key)
                if val is not None:
                    totals[key] += val

        if count > 0:
            ws.append([
                "AVERAGE",
                totals["mem_total"] / count,
                totals["mem_free"] / count,
                totals["mem_used"] / count,
                totals["buff_cache"] / count,
                totals["swap_total"] / count,
                totals["swap_free"] / count,
                totals["swap_used"] / count,
                totals["avail_mem"] / count,
                totals["res"] / count,
                None,
                None,
                totals["virt"] / count,
                totals["res"] / count,
                totals["shr"] / count,
                totals["cpu_pct"] / count,
                totals["mem_pct"] / count
            ])
        else:
            ws.append(["AVERAGE", "(no samples)"])

        ws.append([])

        workbook.save(file_path)
        log_message(f"💾 TOP output saved to Excel: {file_path}")
    except Exception as exc:
        log_message(f"⚠ Could not write TOP output to Excel: {exc}")


def _log_system_command_block(label, samples):
    if not samples:
        return

    header = [
        "Time",
        "MemTotal",
        "MemFree",
        "MemUsed",
        "BuffCache",
        "SwapTotal",
        "SwapFree",
        "SwapUsed",
        "AvailMem",
        "RSS_KB",
        "PSS_KB",
        "PID",
        "VIRT_KiB",
        "RES_KiB",
        "SHR_KiB",
        "CPU_pct",
        "MEM_pct"
    ]

    totals = {
        "mem_total": 0.0, "mem_free": 0.0, "mem_used": 0.0, "buff_cache": 0.0,
        "swap_total": 0.0, "swap_free": 0.0, "swap_used": 0.0, "avail_mem": 0.0,
        "virt": 0.0, "res": 0.0, "shr": 0.0, "cpu_pct": 0.0, "mem_pct": 0.0
    }

    log_message(f"SYSTEM_COMMAND_BLOCK_START|label={label}")
    log_message("SYSTEM_COMMAND_HEADER|" + "|".join(header))

    count = 0
    for sample in samples:
        if not sample:
            continue
        count += 1
        row = [
            sample.get("time"),
            sample.get("mem_total"),
            sample.get("mem_free"),
            sample.get("mem_used"),
            sample.get("buff_cache"),
            sample.get("swap_total"),
            sample.get("swap_free"),
            sample.get("swap_used"),
            sample.get("avail_mem"),
            sample.get("res"),
            None,
            sample.get("pid"),
            sample.get("virt"),
            sample.get("res"),
            sample.get("shr"),
            sample.get("cpu_pct"),
            sample.get("mem_pct")
        ]
        log_message("SYSTEM_COMMAND_ROW|" + "|".join(["" if v is None else str(v) for v in row]))

        for key in totals:
            val = sample.get(key)
            if val is not None:
                totals[key] += val

    if count > 0:
        avg_row = [
            "AVERAGE",
            totals["mem_total"] / count,
            totals["mem_free"] / count,
            totals["mem_used"] / count,
            totals["buff_cache"] / count,
            totals["swap_total"] / count,
            totals["swap_free"] / count,
            totals["swap_used"] / count,
            totals["avail_mem"] / count,
            totals["res"] / count,
            None,
            None,
            totals["virt"] / count,
            totals["res"] / count,
            totals["shr"] / count,
            totals["cpu_pct"] / count,
            totals["mem_pct"] / count
        ]
        log_message("SYSTEM_COMMAND_AVERAGE|" + "|".join(["" if v is None else str(v) for v in avg_row]))

    log_message("SYSTEM_COMMAND_BLOCK_END")


def _resolve_command_by_name(command_name):
    try:
        from controllers.system_commands_controller import SystemCommandsController
        all_cmds = SystemCommandsController.get_all_commands()
        approved = SystemCommandsController.get_approved_commands()

        name = (command_name or '').strip().lower()
        if not name:
            return None

        if name in all_cmds.get('builtin', {}):
            return all_cmds['builtin'][name].get('command')
        if name in all_cmds.get('user_defined', {}):
            return all_cmds['user_defined'][name].get('command')
        if name in approved:
            return approved[name].get('command')
    except Exception as exc:
        log_message(f"⚠ Could not resolve system command '{command_name}': {exc}")
    return None


def execute_system_command(device_ip, port, username, password, iteration=1, device_name="Device",
                           command_name=None, command_text=None, section_label=None,
                           duration_seconds=300, interval_seconds=10, combined_method_name=None):
    timestamp = datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')
    screenshots_list = []
    logs_list = []

    try:
        from services.log_service import LogService
        log_service = LogService()
        method_folder = get_folder_method_name() or combined_method_name
        usb_log_path = _get_system_command_log_path(device_ip, device_name, iteration, method_folder=method_folder)
        log_service.current_usb_log_file = usb_log_path
        log_message(f"💾 USB execution log: {usb_log_path}")
    except Exception as e:
        log_message(f"⚠ Could not set USB log path: {e}")

    log_message("="*80)
    log_message("SYSTEM COMMAND - START")
    log_message("="*80)

    try:
        duration_seconds = int(duration_seconds) if duration_seconds is not None else 300
    except Exception:
        duration_seconds = 300
    try:
        interval_seconds = int(interval_seconds) if interval_seconds is not None else 10
    except Exception:
        interval_seconds = 10

    command_to_run = command_text or _resolve_command_by_name(command_name)
    if not command_to_run:
        msg = "No system command provided or command name not found"
        log_message(f"❌ ERROR: {msg}")
        return {"iteration": iteration, "screenshots": screenshots_list, "logs": logs_list, "success": False, "details": msg}

    log_message(f"📋 Command: {command_to_run}")
    if command_name:
        log_message(f"🏷 Command Name: {command_name}")

    try:
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        ssh.connect(device_ip, port=port, username=username, password=password, timeout=15)
        log_message("✓ Connected to device successfully")

        is_top_command = command_to_run.strip().lower().startswith("top")
        command_output = ""
        exit_status = 0
        top_samples = []

        if is_top_command:
            log_message(f"⏱ Collecting TOP samples for {duration_seconds}s every {interval_seconds}s...")
            end_time = time.time() + max(int(duration_seconds), 1)
            while time.time() < end_time:
                stdin, stdout, stderr = ssh.exec_command(command_to_run, timeout=30)
                sample_output = stdout.read().decode('utf-8', errors='ignore').strip()
                sample_error = stderr.read().decode('utf-8', errors='ignore').strip()
                exit_status = stdout.channel.recv_exit_status()
                if sample_error:
                    log_message(f"⚠ Command Error Output: {sample_error[:200]}")
                command_output = sample_output
                parsed = _parse_top_output(sample_output)
                if parsed:
                    top_samples.append(parsed)
                time.sleep(max(int(interval_seconds), 1))

            log_message(f"✓ TOP collection complete ({len(top_samples)} samples)")
            _log_system_command_block(section_label or command_name or "TOP Capture", top_samples)
            _write_top_output_to_excel(
                device_ip=device_ip,
                device_name=device_name,
                iteration=iteration,
                command=command_to_run,
                samples=top_samples,
                section_label=section_label or command_name,
                combined_method_name=combined_method_name
            )
        else:
            stdin, stdout, stderr = ssh.exec_command(command_to_run, timeout=30)
            command_output = stdout.read().decode('utf-8', errors='ignore').strip()
            command_error = stderr.read().decode('utf-8', errors='ignore').strip()
            exit_status = stdout.channel.recv_exit_status()
            if command_error:
                log_message(f"⚠ Command Error Output: {command_error[:500]}")

        ssh.close()

        success = True if is_top_command else (exit_status == 0)
        details = f"System command executed{' (TOP samples: ' + str(len(top_samples)) + ')' if is_top_command else ''}"

        return {
            "iteration": iteration,
            "screenshots": screenshots_list,
            "logs": logs_list,
            "success": success,
            "details": details,
            "command": command_to_run,
            "output": command_output[:500]
        }

    except paramiko.SSHException as ssh_error:
        log_message(f"❌ SSH Error: {ssh_error}")
        return {
            "iteration": iteration,
            "screenshots": screenshots_list,
            "logs": logs_list,
            "success": False,
            "details": f"SSH connection failed: {str(ssh_error)}"
        }

    except Exception as e:
        import traceback
        log_message(f"❌ Error during System Command execution: {e}")
        log_message(f"Full traceback:\n{traceback.format_exc()}")
        return {
            "iteration": iteration,
            "screenshots": screenshots_list,
            "logs": logs_list,
            "success": False,
            "details": f"Execution error: {str(e)}"
        }
